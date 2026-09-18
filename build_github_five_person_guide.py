from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).with_name("Arbiter Five Person GitHub Workflow.docx")


def cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9D9")


def shade(cell, hex_color):
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), hex_color)
    props.append(node)


def add_text(cell, text, bold=False, color=None, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Aptos"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_border(cell)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        add_text(table.rows[0].cells[i], h, True, "FFFFFF")
        shade(table.rows[0].cells[i], "1F4E78")
        table.rows[0].cells[i].width = Inches(widths[i])
    for n, row in enumerate(rows):
        cells = table.add_row().cells
        for i, v in enumerate(row):
            add_text(cells[i], v)
            cells[i].width = Inches(widths[i])
            if n % 2:
                shade(cells[i], "F3F7FA")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    for r in p.runs:
        r.font.name = "Aptos"
        r._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        r._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        r.font.size = Pt(10)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    p.add_run(text)
    return p


def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.24)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.font.name = "Consolas"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(47, 84, 150)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

for style in ("Normal", "Title", "Heading 1", "Heading 2"):
    doc.styles[style].font.name = "Aptos"
    doc.styles[style]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    doc.styles[style]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    doc.styles[style].font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.add_run("Arbiter Five Person GitHub Workflow")
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
subrun = subtitle.add_run("Exact Windows PowerShell commands for ZIP based individual branches and pull requests")
subrun.font.size = Pt(11)
subrun.font.color.rgb = RGBColor(89, 89, 89)

body(doc, "Use this workflow when each person receives a ZIP file of Arbiter and must push only their assigned code. The safe sequence is: one maintainer publishes the ZIP as main, every contributor downloads that exact GitHub ZIP, each contributor creates one branch from origin/main, edits only owned files, pushes that branch, opens a pull request, and merges only after review.")
body(doc, "Replace every placeholder in angle brackets before running a command. Do not run git add dot. Use the exact git add paths listed for your role so a pull request contains only your work.")

heading(doc, "Repository Values Used Below")
add_table(doc, ["Placeholder", "Replace with"], [
    ["<GITHUB_USER_OR_ORG>", "Your GitHub username or organization, for example arbiter-team"],
    ["<REPOSITORY>", "Repository name, for example Arbiter-Engine"],
    ["<YOUR_NAME>", "Your name, for example Priya Shah"],
    ["<YOUR_EMAIL>", "Email linked to your GitHub account"],
    ["<ZIP_FOLDER>", "Folder produced after extracting the GitHub ZIP"],
], [2.15, 4.6])

heading(doc, "Part One Maintainer Publishes Main From the ZIP")
body(doc, "Only the maintainer performs this step. Create an empty GitHub repository first. Do not add a README, .gitignore, or license during GitHub repository creation. Download or use the project ZIP, extract it, and open PowerShell inside the extracted project folder.")
for line in [
    'cd "C:\\path\\to\\<ZIP_FOLDER>"',
    'git init',
    'git branch -M main',
    'git config user.name "<YOUR_NAME>"',
    'git config user.email "<YOUR_EMAIL>"',
    'git add -A',
    'git commit -m "chore: add Arbiter project baseline"',
    'git remote add origin "https://github.com/<GITHUB_USER_OR_ORG>/<REPOSITORY>.git"',
    'git push -u origin main',
]: code(doc, line)
body(doc, "After the push, open GitHub and confirm that main contains the entire project. Everyone else must now use GitHub Code then Download ZIP for this main branch. This makes every person start from the same committed baseline.")

heading(doc, "Part Two Every Contributor Starts From the Downloaded ZIP")
body(doc, "Each of the five contributors follows these commands once. This local baseline commit is never pushed. It makes the extracted ZIP tracked locally so Git can safely switch to the real remote main branch without deleting untracked files.")
for line in [
    'cd "C:\\path\\to\\<ZIP_FOLDER>"',
    'git init',
    'git branch -M local-zip-baseline',
    'git config user.name "<YOUR_NAME>"',
    'git config user.email "<YOUR_EMAIL>"',
    'git add -A',
    'git commit -m "chore: local ZIP baseline do not push"',
    'git remote add origin "https://github.com/<GITHUB_USER_OR_ORG>/<REPOSITORY>.git"',
    'git fetch origin',
]: code(doc, line)
body(doc, "Next run only the branch command shown in your own row below. Do not push local-zip-baseline. It exists only on your computer.")

heading(doc, "Five Non Overlapping Work Packages")
roles = [
    ["Person 1", "feature/policy-data", "Policy corpus and retrieval", "arbiter/backend/data/policies/; arbiter/backend/data/precedents/; arbiter/backend/agents/retrieval.py; arbiter/backend/stores/policy_store.py; arbiter/backend/init_db.py; tests/test_retrieval.py"],
    ["Person 2", "feature/core-reasoning", "Resolution and adversarial check", "arbiter/backend/agents/resolution.py; arbiter/backend/agents/checker.py; arbiter/backend/orchestrator.py; arbiter/backend/schemas.py; tests/test_resolution.py; tests/test_checker.py"],
    ["Person 3", "feature/decision-intelligence", "Precedent sensitivity and remediation", "arbiter/backend/agents/precedent.py; arbiter/backend/agents/sensitivity.py; arbiter/backend/agents/remediation.py; arbiter/backend/stores/precedent_store.py; tests/test_precedent.py; tests/test_sensitivity.py"],
    ["Person 4", "feature/simulation-graph", "Simulation scanner and policy graph", "arbiter/backend/agents/simulation.py; arbiter/backend/agents/scanner.py; arbiter/backend/stores/graph_store.py; arbiter/backend/data/test_cases/; tests/test_simulation.py; tests/test_scanner.py"],
    ["Person 5", "feature/frontend-identity", "Frontend and sign in experience", "arbiter/frontend/app/; arbiter/frontend/components/; arbiter/frontend/lib/; arbiter/backend/auth.py; arbiter/backend/main.py; tests/test_auth.py"],
]
add_table(doc, ["Person", "Branch", "Work", "Only these paths may be changed"], roles, [0.65, 1.25, 1.25, 3.65])

heading(doc, "Exact Commands for Each Person")
for person, branch, work, paths in roles:
    heading(doc, f"{person} {work}", 2)
    code(doc, f'git switch -c {branch} --track origin/main')
    body(doc, f"Edit only: {paths}")
    code(doc, 'git status')
    body(doc, "Confirm that every changed file belongs to your assigned list. If a file is not assigned to you, restore it before committing.")
    if person == "Person 1":
        add_cmd = 'git add arbiter/backend/data/policies arbiter/backend/data/precedents arbiter/backend/agents/retrieval.py arbiter/backend/stores/policy_store.py arbiter/backend/init_db.py tests/test_retrieval.py'
    elif person == "Person 2":
        add_cmd = 'git add arbiter/backend/agents/resolution.py arbiter/backend/agents/checker.py arbiter/backend/orchestrator.py arbiter/backend/schemas.py tests/test_resolution.py tests/test_checker.py'
    elif person == "Person 3":
        add_cmd = 'git add arbiter/backend/agents/precedent.py arbiter/backend/agents/sensitivity.py arbiter/backend/agents/remediation.py arbiter/backend/stores/precedent_store.py tests/test_precedent.py tests/test_sensitivity.py'
    elif person == "Person 4":
        add_cmd = 'git add arbiter/backend/agents/simulation.py arbiter/backend/agents/scanner.py arbiter/backend/stores/graph_store.py arbiter/backend/data/test_cases tests/test_simulation.py tests/test_scanner.py'
    else:
        add_cmd = 'git add arbiter/frontend/app arbiter/frontend/components arbiter/frontend/lib arbiter/backend/auth.py arbiter/backend/main.py tests/test_auth.py'
    code(doc, add_cmd)
    code(doc, f'git commit -m "feat: {work.lower()}"')
    code(doc, 'git status')
    code(doc, f'git push -u origin {branch}')
    body(doc, f"Open GitHub, choose Compare and pull request for {branch}, set base to main, write a short description, request review, and create the pull request.")

heading(doc, "Pull Request Rules")
for item in [
    "One branch and one pull request per person. Do not push directly to main.",
    "The pull request title must name the work package, for example feat: policy data and retrieval.",
    "The description must list changed files, the feature demonstrated, and tests or build commands run.",
    "Before opening the pull request, run git status. It must not show files outside the assigned work package.",
    "The reviewer checks the file list first. A pull request containing another person’s assigned files is returned for cleanup.",
    "Merge in this order: Person 1, Person 2, Person 3, Person 4, Person 5. After each merge, the next person updates their branch from origin/main before requesting final review.",
]:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(item)
    r.font.size = Pt(9.5)

heading(doc, "Update Your Branch Before Final Review")
body(doc, "If main changed while you were working, run these commands on your feature branch before final review. Resolve conflicts only in your assigned files. If a conflict is in an unassigned file, ask that owner or the maintainer to resolve it.")
for line in [
    'git fetch origin',
    'git merge origin/main',
    'git status',
    'git push',
]: code(doc, line)

heading(doc, "Commands to Avoid")
for item in [
    "Do not use git add dot. It can stage another person’s work or local generated files.",
    "Do not use git push origin main. Only the maintainer merges approved pull requests to main.",
    "Do not use git reset hard, git clean, or force push unless the maintainer explicitly directs it.",
    "Do not upload a new ZIP to GitHub after the initial baseline. Push your branch commits only.",
]:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(item)
    r.font.size = Pt(9.5)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run("Arbiter Five Person GitHub Workflow")
fr.font.size = Pt(8)
fr.font.color.rgb = RGBColor(128, 128, 128)

doc.save(OUT)
print(OUT)
